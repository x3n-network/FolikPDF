"""Generate synthetic interoperability fixtures. Requires python-docx, odfpy, Pillow.
These contain no user data. Run with a temporary Python environment.
"""
from pathlib import Path
from io import BytesIO
import zipfile
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from odf.opendocument import OpenDocumentText
from odf.style import Style, TextProperties, ParagraphProperties, ListLevelProperties
from odf.text import H, P, Span, List, ListItem, ListStyle, ListLevelStyleNumber
from odf.table import Table, TableRow, TableCell, TableColumn
from odf.draw import Frame, Image as ODFImage

out = Path(__file__).parent
image = Image.new('RGB', (160, 100), '#1478bb')
draw = ImageDraw.Draw(image)
draw.rectangle((80, 0, 159, 99), fill='#cc3070')
buffer = BytesIO(); image.save(buffer, format='PNG'); picture = buffer.getvalue()

doc = Document()
doc.add_heading('Engine sample', level=1)
p = doc.add_paragraph('Plain café text, ')
p.add_run('bold text').bold = True
p.add_run(' and italic text').italic = True
r = p.add_run(' with red text.'); r.font.color.rgb = RGBColor(180, 0, 0)
doc.add_paragraph('First bullet', style='List Bullet')
doc.add_paragraph('First numbered item', style='List Number')
doc.add_paragraph('Second numbered item', style='List Number')
doc.add_picture(BytesIO(picture), width=Inches(2))
doc.add_paragraph('AFTER IMAGE MARKER')
t = doc.add_table(rows=2, cols=2)
t.cell(0,0).text='Column A'; t.cell(0,1).text='Column B'
t.cell(1,0).text='\n'.join(f'Cell line {i:03d}' for i in range(160))
t.cell(1,1).text='Neighbor cell'
merged=t.add_row().cells[0].merge(t.rows[-1].cells[1])
merged.text='Merged cell with image'
merged.paragraphs[0].add_run().add_picture(BytesIO(picture), width=Inches(1))
doc.add_paragraph('AFTER TABLE MARKER')
doc.add_page_break()
doc.add_paragraph('FINAL DOCX MARKER')
doc.sections[0].header.paragraphs[0].text='Header omitted with warning'
doc.save(out/'OfficeEngine.docx')

odt = OpenDocumentText()
heading = Style(name='Heading1', family='paragraph')
heading.addElement(TextProperties(fontsize='20pt', fontweight='bold'))
heading.addElement(ParagraphProperties(keepwithnext='always'))
odt.styles.addElement(heading)
bold = Style(name='Bold', family='text'); bold.addElement(TextProperties(fontweight='bold')); odt.styles.addElement(bold)
odt.text.addElement(H(outlinelevel=1, stylename=heading, text='Engine sample'))
p = P(text='Plain café text, '); p.addElement(Span(stylename=bold, text='bold text')); odt.text.addElement(p)
ls = ListStyle(name='Numbers')
ls.addElement(ListLevelStyleNumber(level=1, numformat='1', numsuffix='.', startvalue=3))
odt.styles.addElement(ls)
lst=List(stylename=ls)
for text in ['First numbered item', 'Second numbered item']:
 item=ListItem(); item.addElement(P(text=text)); lst.addElement(item)
odt.text.addElement(lst)
frame=Frame(width='2in', height='1.25in', anchortype='as-char')
frame.addElement(ODFImage(href=odt.addPicture('Pictures/sample.png', mediatype='image/png', content=picture), type='simple', show='embed', actuate='onLoad'))
p=P(); p.addElement(frame); odt.text.addElement(p)
odt.text.addElement(P(text='AFTER IMAGE MARKER'))
table=Table(name='Sample'); table.addElement(TableColumn(numbercolumnsrepeated=2))
for values in [('Column A','Column B'), ('long','Neighbor cell')]:
 row=TableRow()
 for value in values:
  cell=TableCell(valuetype='string')
  if value=='long':
   for i in range(160): cell.addElement(P(text=f'Cell line {i:03d}'))
  else: cell.addElement(P(text=value))
  row.addElement(cell)
 table.addElement(row)
row=TableRow()
cell=TableCell(valuetype='string', numbercolumnsspanned=2)
cell.addElement(P(text='Merged cell with image'))
frame=Frame(width='1in', height='0.625in', anchortype='as-char')
frame.addElement(ODFImage(href='Pictures/sample.png', type='simple', show='embed', actuate='onLoad'))
p=P(); p.addElement(frame); cell.addElement(p); row.addElement(cell)
from odf.table import CoveredTableCell
row.addElement(CoveredTableCell()); table.addElement(row)
odt.text.addElement(table)
odt.text.addElement(P(text='AFTER TABLE MARKER'))
odt.text.addElement(P(text='FINAL ODT MARKER'))
odt.save(str(out/'OfficeEngine.odt'))

# Minimal independent archives for package-integrity tests.
for method, name in [(zipfile.ZIP_STORED,'OfficeStored.zip'), (zipfile.ZIP_DEFLATED,'OfficeDeflated.zip')]:
 with zipfile.ZipFile(out/name,'w',compression=method) as z: z.writestr('content.xml','known archive content')
# Compact preview fixtures for UI launch arguments; separate from pagination stress cases.
for ext in ['docx','odt']:
 source=out/f'OfficeEngine.{ext}'
 with zipfile.ZipFile(source) as z, zipfile.ZipFile(out/f'OfficePreview.{ext}','w',compression=zipfile.ZIP_DEFLATED) as target:
  for info in z.infolist():
   data=z.read(info.filename)
   if info.filename in ['word/document.xml','content.xml']:
    import re
    text=data.decode('utf-8')
    if ext=='docx':
     text=re.sub(r'<w:t>Cell line \d+</w:t>', '<w:t>Short cell</w:t>', text, count=1)
     text=re.sub(r'<w:br/><w:t>Cell line \d+</w:t>', '', text)
    else:
     text=re.sub(r'<text:p>Cell line \d+</text:p>', '', text)
    data=text.encode('utf-8')
   target.writestr(info.filename,data)
print('Generated synthetic DOCX/ODT and archive fixtures.')
